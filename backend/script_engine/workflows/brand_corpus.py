"""
Magppie Brand Corpus — the base of research.

Everything the engine may assert *about Magppie* must trace to a document uploaded
here. Model-weight recall is still allowed for general history and science, but it
is marked `general-knowledge` and can never become a product claim or back a
`verified-stat` hook. `corpus_gate.py` enforces that; this module stores the truth.

Three deliberate design choices:

1. **Uploads are not auto-trusted.** Extraction from a PDF is imperfect, and a
   mis-parsed dimension becomes a false product specification. Raw text is stored
   with provenance and marked `pending` until a human promotes it.
2. **Selection, not concatenation.** A real brand book is far larger than a prompt
   budget. `select_for_topic` returns only the relevant slices.
3. **The original is never edited.** `_uploads/<id>/` keeps the source file and the
   raw extraction verbatim, so any claim can be traced back to its page.
"""

import hashlib
import json
import logging
import re
import shutil
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

logger = logging.getLogger("MagppieEngine.BrandCorpus")

UPLOADS_DIRNAME = "_uploads"
MANIFEST_FILENAME = "manifest.json"

# The seven categories that make up Magppie truth. Order matters: it is the order
# they are offered in the UI and assembled into an agent prompt.
CATEGORIES = (
    "brand",
    "products",
    "usps",
    "approved-claims",
    "prohibited-claims",
    "visual-identity",
    "references",
)

CATEGORY_LABELS = {
    "brand": "Brand — history, positioning, pillars, voice",
    "products": "Products — materials, finishes, dimensions, certifications",
    "usps": "USPs — differentiators and the evidence behind them",
    "approved-claims": "Approved claims — what marketing may state",
    "prohibited-claims": "Prohibited claims — what must never be said",
    "visual-identity": "Visual identity — palette, materials, photography style",
    "references": "References — standards, test reports, source hierarchy",
}

# Where a claim came from. This is the field the corpus gate reasons about.
ORIGIN_MAGPPIE = "magppie-corpus"
ORIGIN_GENERAL = "general-knowledge"

SUPPORTED_SUFFIXES = {".pdf", ".md", ".markdown", ".txt", ".docx", ".rst", ".csv"}

MAX_UPLOAD_BYTES = 25 * 1024 * 1024

_WORD_RE = re.compile(r"[a-z0-9][a-z0-9\-]{2,}")

_STOPWORDS = frozenset("""
the and for with from that this those these into onto are was were been being its
has have had not but you your our their they them there here what when where which
who whom how why all any can may will would should could about over under more most
""".split())


class CorpusError(Exception):
    """Raised when an upload cannot be accepted or read."""


# ---------------------------------------------------------------- extraction

def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise CorpusError("pypdf is required to read PDF uploads") from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise CorpusError(f"Could not open this PDF: {exc}") from exc
    pages = []
    for number, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        # Page markers survive into the stored text so a claim can cite its page.
        pages.append(f"\n[page {number}]\n{text.strip()}")
    return "\n".join(pages).strip()


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise CorpusError("python-docx is required to read .docx uploads") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise CorpusError(f"Could not open this Word document: {exc}") from exc
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


def extract(path: Path) -> str:
    """Plain text from a supported document. Raises CorpusError on anything else."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    elif suffix in SUPPORTED_SUFFIXES:
        text = _extract_text(path)
    else:
        raise CorpusError(
            f"Unsupported file type '{suffix}'. Supported: "
            + ", ".join(sorted(SUPPORTED_SUFFIXES))
        )
    if not text.strip():
        raise CorpusError(
            "No text could be read from this file. A scanned PDF needs OCR before "
            "it can be used as evidence."
        )
    return text


# ---------------------------------------------------------------- the corpus

class BrandCorpus:
    """Stores uploaded Magppie source documents and serves topic-relevant slices."""

    def __init__(self, knowledge_dir: Path):
        self.root = Path(knowledge_dir)
        self.uploads_dir = self.root / UPLOADS_DIRNAME

    # ---------- manifest ----------

    @property
    def manifest_path(self) -> Path:
        return self.uploads_dir / MANIFEST_FILENAME

    def _read_manifest(self) -> List[Dict[str, Any]]:
        if not self.manifest_path.exists():
            return []
        try:
            data = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            logger.warning("Corpus manifest unreadable; treating as empty.")
            return []
        return data if isinstance(data, list) else []

    def _write_manifest(self, entries: List[Dict[str, Any]]) -> None:
        self.uploads_dir.mkdir(parents=True, exist_ok=True)
        self.manifest_path.write_text(
            json.dumps(entries, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    # ---------- ingestion ----------

    def add_document(
        self,
        source_path: Path,
        category: str,
        original_filename: Optional[str] = None,
        notes: str = "",
    ) -> Dict[str, Any]:
        """
        Store a document, extract its text, and register it as `pending`.

        Pending means: readable, but not yet promoted to Magppie truth. Extraction
        is imperfect and a mis-parsed specification would become a false product
        claim, so a human confirms before agents may cite it.
        """
        if category not in CATEGORIES:
            raise CorpusError(
                f"Unknown category '{category}'. Expected one of: {', '.join(CATEGORIES)}"
            )
        source_path = Path(source_path)
        if not source_path.exists():
            raise CorpusError(f"File not found: {source_path}")

        size = source_path.stat().st_size
        if size == 0:
            raise CorpusError("That file is empty.")
        if size > MAX_UPLOAD_BYTES:
            raise CorpusError(
                f"File is {size / 1024 / 1024:.1f} MB; the limit is "
                f"{MAX_UPLOAD_BYTES // 1024 // 1024} MB."
            )

        filename = original_filename or source_path.name
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise CorpusError(
                f"Unsupported file type '{suffix}'. Supported: "
                + ", ".join(sorted(SUPPORTED_SUFFIXES))
            )

        digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
        existing = next((e for e in self._read_manifest() if e.get("sha256") == digest), None)
        if existing:
            raise CorpusError(
                f"This document is already in the corpus as '{existing['filename']}' "
                f"({existing['doc_id']}). Remove it first to replace it."
            )

        doc_id = f"DOC-{digest[:10].upper()}"
        folder = self.uploads_dir / doc_id
        folder.mkdir(parents=True, exist_ok=True)

        stored = folder / f"original{suffix}"
        shutil.copy2(source_path, stored)

        try:
            text = extract(stored)
        except CorpusError:
            shutil.rmtree(folder, ignore_errors=True)
            raise

        (folder / "extracted.txt").write_text(text, encoding="utf-8")

        entry = {
            "doc_id": doc_id,
            "filename": filename,
            "category": category,
            "status": "pending",          # pending → approved (human promotes)
            "sha256": digest,
            "bytes": size,
            "characters": len(text),
            "estimated_tokens": len(text) // 4,
            "notes": notes,
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "approved_at": None,
        }
        entries = self._read_manifest()
        entries.insert(0, entry)
        self._write_manifest(entries)
        logger.info("[corpus] stored %s (%s, %d chars) as pending", doc_id, category, len(text))
        return entry

    def set_status(self, doc_id: str, status: str) -> Dict[str, Any]:
        """Promote a document to `approved`, or send it back to `pending`."""
        if status not in ("pending", "approved"):
            raise CorpusError("Status must be 'pending' or 'approved'.")
        entries = self._read_manifest()
        for entry in entries:
            if entry.get("doc_id") == doc_id:
                entry["status"] = status
                entry["approved_at"] = (
                    datetime.now(timezone.utc).isoformat() if status == "approved" else None
                )
                self._write_manifest(entries)
                return entry
        raise CorpusError(f"No such document: {doc_id}")

    def remove(self, doc_id: str) -> None:
        entries = self._read_manifest()
        remaining = [e for e in entries if e.get("doc_id") != doc_id]
        if len(remaining) == len(entries):
            raise CorpusError(f"No such document: {doc_id}")
        shutil.rmtree(self.uploads_dir / doc_id, ignore_errors=True)
        self._write_manifest(remaining)

    # ---------- reads ----------

    def documents(self, category: Optional[str] = None,
                  status: Optional[str] = None) -> List[Dict[str, Any]]:
        entries = self._read_manifest()
        if category:
            entries = [e for e in entries if e.get("category") == category]
        if status:
            entries = [e for e in entries if e.get("status") == status]
        return entries

    def document_text(self, doc_id: str) -> str:
        path = self.uploads_dir / doc_id / "extracted.txt"
        if not path.exists():
            raise CorpusError(f"No extracted text for {doc_id}")
        return path.read_text(encoding="utf-8", errors="replace")

    def is_populated(self) -> bool:
        """True when at least one approved document exists."""
        return bool(self.documents(status="approved"))

    def stats(self) -> Dict[str, Any]:
        entries = self._read_manifest()
        approved = [e for e in entries if e.get("status") == "approved"]
        by_category = {c: 0 for c in CATEGORIES}
        for entry in approved:
            by_category[entry.get("category", "brand")] = \
                by_category.get(entry.get("category", "brand"), 0) + 1
        return {
            "documents": len(entries),
            "approved": len(approved),
            "pending": len(entries) - len(approved),
            "approved_characters": sum(e.get("characters", 0) for e in approved),
            "by_category": by_category,
            "populated": bool(approved),
        }

    # ---------- selection ----------

    @staticmethod
    def _keywords(text: str) -> List[str]:
        folded = unicodedata.normalize("NFKD", (text or "").lower())
        return [w for w in _WORD_RE.findall(folded) if w not in _STOPWORDS]

    @staticmethod
    def _chunk(text: str, target_chars: int = 1200) -> List[str]:
        """Split on blank lines, then pack into chunks that stay near target size."""
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        chunks, current = [], ""
        for paragraph in paragraphs:
            if current and len(current) + len(paragraph) + 2 > target_chars:
                chunks.append(current)
                current = paragraph
            else:
                current = f"{current}\n\n{paragraph}" if current else paragraph
        if current:
            chunks.append(current)
        return chunks

    def select_for_topic(
        self,
        topic: str,
        objective: str = "",
        product: str = "",
        char_budget: int = 9000,
        categories: Optional[Iterable[str]] = None,
    ) -> Dict[str, Any]:
        """
        The slice of the corpus relevant to one run.

        A full brand book cannot ride in every prompt — Agent 1's prompt is already
        ~8k tokens and free-tier throughput is the bottleneck. Chunks are scored by
        keyword overlap with the brief and packed until the budget is spent.

        Prohibited claims are **always** included regardless of score: the guardrail
        must never be dropped for being off-topic.
        """
        wanted = set(categories or CATEGORIES)
        query = set(self._keywords(f"{topic} {objective} {product}"))
        selected: List[Dict[str, Any]] = []
        always: List[Dict[str, Any]] = []

        for entry in self.documents(status="approved"):
            if entry.get("category") not in wanted:
                continue
            try:
                text = self.document_text(entry["doc_id"])
            except CorpusError:
                continue
            for index, chunk in enumerate(self._chunk(text)):
                record = {
                    "doc_id": entry["doc_id"],
                    "filename": entry["filename"],
                    "category": entry["category"],
                    "chunk": index,
                    "text": chunk,
                }
                if entry["category"] == "prohibited-claims":
                    always.append(record)
                    continue
                overlap = query & set(self._keywords(chunk))
                record["score"] = len(overlap)
                record["matched"] = sorted(overlap)[:8]
                selected.append(record)

        selected.sort(key=lambda r: (-r["score"], r["doc_id"], r["chunk"]))

        packed: List[Dict[str, Any]] = []
        used = 0
        for record in always:           # guardrails first, unconditionally
            packed.append(record)
            used += len(record["text"])
        for record in selected:
            if used + len(record["text"]) > char_budget:
                continue
            packed.append(record)
            used += len(record["text"])

        return {
            "chunks": packed,
            "characters": used,
            "documents_used": sorted({r["doc_id"] for r in packed}),
            "budget": char_budget,
        }

    def prompt_block(self, topic: str, objective: str = "", product: str = "",
                     char_budget: int = 9000) -> str:
        """The corpus slice formatted for an agent's system prompt."""
        chosen = self.select_for_topic(topic, objective, product, char_budget)
        if not chosen["chunks"]:
            return (
                "# MAGPPIE BRAND CORPUS\n"
                "No approved Magppie source documents have been uploaded yet.\n"
                "You therefore have NO verified Magppie-specific evidence. Do not assert "
                "any Magppie material, specification, certification, dimension or figure. "
                "General historical and scientific context may still be used, but it must "
                f"be marked evidence_origin='{ORIGIN_GENERAL}' and verification_status="
                "'UNVERIFIED'.\n"
            )
        lines = [
            "# MAGPPIE BRAND CORPUS (the only valid source of Magppie truth)",
            "",
            f"Any claim ABOUT MAGPPIE must trace to one of these documents and carry "
            f"evidence_origin='{ORIGIN_MAGPPIE}' plus its source_doc_id.",
            f"General history or science recalled from your own knowledge is permitted, "
            f"but must carry evidence_origin='{ORIGIN_GENERAL}' and "
            f"verification_status='UNVERIFIED'. It may never become a product claim "
            f"or support a 'verified-stat' hook.",
            "",
        ]
        for record in chosen["chunks"]:
            lines.append(
                f"--- [{record['doc_id']}] {record['category']} · {record['filename']} ---"
            )
            lines.append(record["text"])
            lines.append("")
        return "\n".join(lines)
