"""Company knowledge base: loading the user's files and searching them (the RAG part)."""

import csv
import io
import logging
import re
from pathlib import Path

from . import db
from .config import get_settings
from .embed import embed, embed_one

log = logging.getLogger("mai.kb")

CHUNK_CHARS = 1200
OVERLAP = 150


def chunk_markdown(text: str) -> list[tuple[str, str]]:
    """Split markdown into (section heading path, chunk) pairs, keeping headings as context."""
    chunks: list[tuple[str, str]] = []
    heads: list[str] = []
    buf: list[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if not body:
            return
        section = " > ".join(heads)
        start = 0
        while start < len(body):
            piece = body[start:start + CHUNK_CHARS]
            chunks.append((section, piece))
            if start + CHUNK_CHARS >= len(body):
                break
            start += CHUNK_CHARS - OVERLAP

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush()
            buf = []
            level = len(m.group(1))
            heads = heads[: level - 1] + [m.group(2).strip()]
        else:
            buf.append(line)
    flush()
    return chunks


def read_script_bank(path: Path) -> list[str]:
    """90 script ideas from .md / .txt / .csv / .xlsx — one idea per item."""
    suffix = path.suffix.lower()
    if suffix in (".csv",):
        rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8", errors="ignore"))))
        return [" | ".join(c for c in r if c.strip()) for r in rows[1:] if any(c.strip() for c in r)]
    if suffix in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        items = []
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = [str(c) for c in row if c not in (None, "")]
                if i and cells:
                    items.append(" | ".join(cells))
        return items
    text = path.read_text(encoding="utf-8", errors="ignore")
    # numbered or bulleted items, or blocks separated by blank lines / headings
    blocks = re.split(r"\n\s*\n|\n(?=#{1,6}\s)|\n(?=\s*(?:\d+[.)]|[-*])\s)", text)
    return [b.strip() for b in blocks if len(b.strip()) > 20]


async def _insert(doc: str, pairs: list[tuple[str, str]], meta: dict | None = None):
    if not pairs:
        return 0
    vectors = await embed([f"{s}\n{c}" if s else c for s, c in pairs])
    async with db.conn() as c:
        async with c.cursor() as cur:
            await cur.executemany(
                "insert into kb_chunks (doc, section, content, embedding, meta) values (%s,%s,%s,%s,%s)",
                [(doc, s, ch, v, db.Jsonb(meta or {})) for (s, ch), v in zip(pairs, vectors)],
            )
    return len(pairs)


async def load_inputs() -> dict:
    """(Re)load the core KB, brand guide and script bank from the inputs folder."""
    s = get_settings()
    d = s.inputs_dir
    counts = {}
    await db.execute("delete from kb_chunks where doc in ('core','brand','script_bank','internal')")

    kb_files = sorted(p for p in d.glob("*.md") if p.stem.lower().startswith(("company-kb", "company_kb")) or p.stem.lower() == "kb")
    for p in kb_files:
        # company-kb-internal*.md is for the team only: MAGI may tell it, content agents never see it
        doc = "internal" if "internal" in p.stem.lower() else "core"
        counts[doc] = counts.get(doc, 0) + await _insert(
            doc, chunk_markdown(p.read_text(encoding="utf-8", errors="ignore")), {"file": p.name})

    from . import inputs
    brand_text = inputs.brand_guide()
    if brand_text:
        counts["brand"] = await _insert("brand", chunk_markdown(brand_text), {"file": "brand-guide.md"})

    for p in sorted(d.glob("script-bank.*")) + sorted(d.glob("script_bank.*")):
        ideas = read_script_bank(p)
        counts["script_bank"] = counts.get("script_bank", 0) + await _insert(
            "script_bank", [("Script idea", i) for i in ideas], {"file": p.name})
    bank = d / "prompt-bank.json"
    if bank.exists():                       # earlier Higgsfield prompts the team wrote: style reference
        import json
        items = json.loads(bank.read_text(encoding="utf-8")).get("items", [])
        counts["script_bank"] = counts.get("script_bank", 0) + await _insert(
            "script_bank", [(f"Earlier Higgsfield {i['type']} prompt", i["prompt"][:1500]) for i in items],
            {"file": bank.name})
    log.info("KB loaded: %s", counts)
    return counts


async def search(query: str, k: int = 8, docs: list[str] | None = None) -> list[dict]:
    q = await embed_one(query)
    return await db.fetch("select * from match_kb(%s, %s, %s)", q, k, docs)


def as_context(rows: list[dict]) -> str:
    return "\n\n".join(
        f"[KB-{r['id']} · {r['doc']}{' · ' + r['section'] if r.get('section') else ''}]\n{r['content']}"
        for r in rows)


async def add_learned(content: str, meta: dict) -> None:
    await _insert("learned", [("Learned", content)], meta)


async def add_manager_knowledge(content: str) -> None:
    await _insert("manager", [("Added by manager", content)], {"source": "manager"})
